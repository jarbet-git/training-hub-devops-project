from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session

from app.api.deps import require_hr_or_admin
from app.core.database import get_db
from app.models.area import Area
from app.models.business_need import BusinessNeed
from app.models.cost_center import CostCenter
from app.models.enums import FormStatus
from app.models.export_preset import ExportPreset
from app.models.form import Form
from app.models.form_item import FormItem
from app.models.training_category import TrainingCategory
from app.models.training_name import TrainingName
from app.reports.export_columns import (
    DEFAULT_COLUMNS_V1,
    header_of,
    headers_for,
    normalize_columns,
    summary_headers,
    summary_row,
)

router = APIRouter(prefix="/admin/reports", tags=["admin-reports"], dependencies=[Depends(require_hr_or_admin)])


def _dict_name(obj: Any, lang: str) -> str:
    """Prefer name in requested lang, then fallback to the other lang, then code."""
    if obj is None:
        return "—"

    name_pl = getattr(obj, "name_pl", None)
    name_en = getattr(obj, "name_en", None)
    code = getattr(obj, "code", None)

    if (lang or "pl") == "en":
        return (name_en or name_pl or code or "—").strip() if isinstance((name_en or name_pl or code), str) else (
            name_en or name_pl or code or "—"
        )

    # default: pl
    return (name_pl or name_en or code or "—").strip() if isinstance((name_pl or name_en or code), str) else (
        name_pl or name_en or code or "—"
    )


def _code_name(obj: Any, lang: str) -> str:
    if obj is None:
        return "—"
    code = getattr(obj, "code", None)
    name = _dict_name(obj, lang)
    if code and name and name != code:
        return f"{code} — {name}"
    return str(code or name or "—")


def _priority_label(priority: Optional[str], lang: str) -> str:
    p = (priority or "").upper()
    if lang == "en":
        return {"LOW": "Low", "MEDIUM": "Medium", "HIGH": "High"}.get(p, p or "—")
    return {"LOW": "Niski", "MEDIUM": "Średni", "HIGH": "Wysoki"}.get(p, p or "—")


def _status_label(status: Optional[str], lang: str) -> str:
    s = (status or "").upper()
    if lang == "en":
        return {
            "DRAFT": "Draft",
            "MANAGER_REVIEW": "Manager review",
            "HR_REVIEW": "Sent to HR",
            "REPLIED": "Replied",
        }.get(s, s or "—")
    return {
        "DRAFT": "Szkic",
        "MANAGER_REVIEW": "Weryfikacja kierownika",
        "HR_REVIEW": "Wysłano do HR",
        "REPLIED": "Zakończony",
    }.get(s, s or "—")


def _decision_label(decision: Optional[str], lang: str) -> str:
    d = (decision or "").upper()
    if not d:
        return "—"
    if lang == "en":
        return {"APPROVED": "Approved", "PARTIAL": "Approved", "REJECTED": "Rejected"}.get(d, d)
    return {"APPROVED": "Zaakceptowano", "PARTIAL": "Zaakceptowano", "REJECTED": "Odrzucono"}.get(d, d)


def _dt_range_from_dates(date_from: Optional[date], date_to: Optional[date]) -> tuple[Optional[datetime], Optional[datetime]]:
    """Return [start, end) datetimes in UTC-naive terms."""
    start_dt: Optional[datetime] = None
    end_dt: Optional[datetime] = None

    if date_from:
        start_dt = datetime.combine(date_from, time.min)
    if date_to:
        end_dt = datetime.combine(date_to + timedelta(days=1), time.min)

    return start_dt, end_dt


def _write_xlsx(
    lang: str,
    rows: List[Dict[str, Any]],
    headers: List[str],
    totals_cost: float,
    totals_hours: float,
    per_form_summary: List[Dict[str, Any]],
) -> bytes:
    lang = (lang or "pl").lower()
    if lang not in {"pl", "en"}:
        lang = "pl"

    wb = Workbook()
    ws = wb.active
    ws.title = "Pozycje" if lang == "pl" else "Items"

    header_font = Font(bold=True)

    # header
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    # rows
    for r in rows:
        ws.append([r.get(h) for h in headers])

    # totals row
    if rows:
        ws.append([None] * len(headers))
        totals_row_idx = ws.max_row + 1
        totals_row = [None] * len(headers)
        totals_row[0] = "SUMA" if lang == "pl" else "TOTAL"

        def col_index(name: str) -> int:
            return headers.index(name)

        # place totals if those columns exist
        total_cost_h = header_of("total_cost", lang)
        total_hours_h = header_of("total_hours", lang)
        if total_cost_h in headers:
            totals_row[col_index(total_cost_h)] = float(totals_cost)
        if total_hours_h in headers:
            totals_row[col_index(total_hours_h)] = float(totals_hours)

        ws.append(totals_row)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=totals_row_idx, column=col_idx)
            c.font = header_font

    # formats
    def set_col_format(header_name: str, number_format: str):
        if header_name not in headers:
            return
        idx = headers.index(header_name) + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=idx)
            if isinstance(cell.value, (int, float)):
                cell.number_format = number_format

    set_col_format(header_of("cost_per_person", lang), "#,##0.00")
    set_col_format(header_of("total_cost", lang), "#,##0.00")
    set_col_format(header_of("hr_budget_total", lang), "#,##0.00")
    set_col_format(header_of("hours_per_person", lang), "0.0")
    set_col_format(header_of("total_hours", lang), "0.0")

    # layout
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # column widths
    for i, h in enumerate(headers, start=1):
        max_len = len(str(h))
        for r in range(2, min(ws.max_row, 200) + 1):
            v = ws.cell(row=r, column=i).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, max_len + 2), 48)

    # ---- 2nd sheet: summary per form ----
    ws2 = wb.create_sheet("Wnioski" if lang == "pl" else "Forms")
    if per_form_summary:
        sum_headers = list(per_form_summary[0].keys())
        ws2.append(sum_headers)
        for col_idx in range(1, len(sum_headers) + 1):
            c = ws2.cell(row=1, column=col_idx)
            c.font = header_font
            c.alignment = Alignment(vertical="center")

        for row in per_form_summary:
            ws2.append([row.get(h) for h in sum_headers])

        ws2.freeze_panes = "A2"
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(sum_headers))}1"

        def set_ws2_format(header_name: str, number_format: str):
            if header_name not in sum_headers:
                return
            idx = sum_headers.index(header_name) + 1
            for r in range(2, ws2.max_row + 1):
                cell = ws2.cell(row=r, column=idx)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = number_format

        # use translated headers
        sh = summary_headers(lang)
        # sum_cost, sum_hours, hr_budget_total are at fixed indices in helper
        set_ws2_format(sh[5], "#,##0.00")
        set_ws2_format(sh[8], "#,##0.00")
        set_ws2_format(sh[6], "0.0")

        for i, h in enumerate(sum_headers, start=1):
            max_len = len(str(h))
            for r in range(2, min(ws2.max_row, 200) + 1):
                v = ws2.cell(row=r, column=i).value
                if v is None:
                    continue
                max_len = max(max_len, len(str(v)))
            ws2.column_dimensions[get_column_letter(i)].width = min(max(10, max_len + 2), 48)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _parse_list_param(values: Optional[List[str]]) -> List[str]:
    if not values:
        return []
    if len(values) == 1 and "," in values[0]:
        return [v.strip() for v in values[0].split(",") if v.strip()]
    return [v.strip() for v in values if v and v.strip()]


@router.get("/forms-export")
def export_forms(
    db: Session = Depends(get_db),
    format: str = Query("xlsx", description="xlsx|csv"),
    lang: str = Query("pl", description="pl|en"),
    # PRO: columns/presets
    preset_id: Optional[int] = Query(None, description="Export preset id"),
    columns: Optional[List[str]] = Query(None, description="Repeat or comma-separate column keys"),
    # filters
    status: Optional[str] = Query(None, description="Form status (DRAFT/MANAGER_REVIEW/HR_REVIEW/REPLIED)"),
    statuses: Optional[List[str]] = Query(None, description="Multi-select statuses"),
    # backward compat
    status_value: Optional[str] = Query(None, deprecated=True),
    area_id: Optional[int] = Query(None),
    area_ids: Optional[List[int]] = Query(None, description="Multi-select area ids"),
    created_by_user_id: Optional[int] = Query(None),
    hr_decision: Optional[str] = Query(None, description="HR decision (APPROVED/REJECTED/NONE)"),
    hr_decisions: Optional[List[str]] = Query(None, description="Multi-select HR decisions"),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    # backward compat
    created_from: Optional[datetime] = Query(None, deprecated=True),
    created_to: Optional[datetime] = Query(None, deprecated=True),
) -> StreamingResponse:
    fmt = (format or "xlsx").lower().strip()
    if fmt not in {"xlsx", "csv"}:
        raise HTTPException(status_code=400, detail="Invalid format. Use xlsx or csv.")

    lang = (lang or "pl").lower().strip()
    if lang not in {"pl", "en"}:
        lang = "pl"

    # ---- columns / preset ----
    cols_raw = _parse_list_param(columns)
    columns_keys: List[str]
    if cols_raw:
        try:
            columns_keys = normalize_columns(cols_raw)
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))
    elif preset_id is not None:
        preset = db.get(ExportPreset, preset_id)
        if not preset:
            raise HTTPException(status_code=404, detail="Preset not found")
        try:
            columns_keys = normalize_columns(list(preset.columns or []))
        except ValueError:
            columns_keys = list(DEFAULT_COLUMNS_V1)
    else:
        columns_keys = list(DEFAULT_COLUMNS_V1)

    headers = headers_for(columns_keys, lang)

    # ---- status filter (single + multi) ----
    status_list = _parse_list_param(statuses)
    if not status_list:
        status_raw = (status or status_value or "").strip().upper()
        status_list = [status_raw] if status_raw and status_raw != "ALL" else []

    status_enums: List[FormStatus] = []
    for s in status_list:
        if not s or s == "ALL":
            continue
        try:
            status_enums.append(FormStatus(s))
        except Exception:
            raise HTTPException(status_code=400, detail=f"Invalid status: {s}")

    # ---- decision filter (single + multi) ----
    decision_list = _parse_list_param(hr_decisions)
    if not decision_list:
        d = (hr_decision or "").strip().upper()
        decision_list = [d] if d and d != "ALL" else []

    # date range
    start_dt: Optional[datetime] = created_from
    end_dt: Optional[datetime] = created_to

    if not start_dt and not end_dt:
        s, e = _dt_range_from_dates(date_from, date_to)
        start_dt, end_dt = s, e

    stmt = (
        select(
            Form,
            FormItem,
            Area,
            CostCenter,
            TrainingCategory,
            TrainingName,
            BusinessNeed,
        )
        .join(FormItem, FormItem.form_id == Form.id)
        .join(Area, Area.id == Form.area_id)
        .join(CostCenter, CostCenter.id == FormItem.cost_center_id)
        .join(TrainingCategory, TrainingCategory.id == FormItem.training_category_id)
        .join(TrainingName, TrainingName.id == FormItem.training_name_id)
        .join(BusinessNeed, BusinessNeed.id == FormItem.business_need_id)
    )

    filters = []

    if status_enums:
        filters.append(Form.status.in_(status_enums))

    # area (single + multi)
    if area_ids:
        filters.append(Form.area_id.in_(area_ids))
    elif area_id is not None:
        filters.append(Form.area_id == area_id)

    if created_by_user_id is not None:
        filters.append(Form.created_by_user_id == created_by_user_id)

    if start_dt is not None:
        filters.append(Form.created_at >= start_dt)
    if end_dt is not None:
        filters.append(Form.created_at < end_dt)

    if decision_list:
        # support NONE meaning NULL
        wants_none = any(d in {"NONE", "NULL", "EMPTY"} for d in decision_list)
        real = [d for d in decision_list if d and d not in {"ALL", "NONE", "NULL", "EMPTY"}]
        parts = []
        if real:
            parts.append(Form.hr_decision.in_(real))
        if wants_none:
            parts.append(Form.hr_decision.is_(None))
        if parts:
            filters.append(or_(*parts))

    if filters:
        stmt = stmt.where(and_(*filters))

    stmt = stmt.order_by(Form.id.desc(), FormItem.id.asc())

    results = db.execute(stmt).all()

    # rows
    rows: List[Dict[str, Any]] = []
    totals_cost = 0.0
    totals_hours = 0.0

    # per-form summary helpers
    form_totals_cost: Dict[int, float] = defaultdict(float)
    form_totals_hours: Dict[int, float] = defaultdict(float)
    form_cost_centers: Dict[int, set[str]] = defaultdict(set)
    form_meta: Dict[int, Dict[str, Any]] = {}

    for form, it, area, cc, cat, tn, bn in results:
        employees = int(it.employees_count or 0)
        cost_pp = float(it.estimated_cost_per_person or 0)
        hours_pp = float(it.estimated_hours_per_person or 0)

        item_total_cost = employees * cost_pp
        item_total_hours = employees * hours_pp

        totals_cost += item_total_cost
        totals_hours += item_total_hours

        form_totals_cost[form.id] += item_total_cost
        form_totals_hours[form.id] += item_total_hours
        form_cost_centers[form.id].add(getattr(cc, "code", "") or "")

        if form.id not in form_meta:
            form_meta[form.id] = {
                "form_id": form.id,
                "area": _code_name(area, lang),
                "status": _status_label(getattr(form.status, "value", str(form.status)), lang),
                "created_at": form.created_at.strftime("%Y-%m-%d %H:%M") if getattr(form, "created_at", None) else "—",
                "hr_decision": _decision_label(form.hr_decision, lang),
                "hr_budget_total": float(form.hr_budget_total) if form.hr_budget_total is not None else None,
                "hr_comment": (form.hr_comment or "").strip() or None,
            }

        values_by_key: Dict[str, Any] = {
            "area": _code_name(area, lang),
            "mpk": _code_name(cc, lang),
            "training": _dict_name(tn, lang),
            "category": _dict_name(cat, lang),
            "business_need": _dict_name(bn, lang),
            "employees_count": employees,
            "cost_per_person": cost_pp,
            "total_cost": item_total_cost,
            "hours_per_person": hours_pp,
            "total_hours": item_total_hours,
            "quarter": (it.quarter or "").upper() or "—",
            "priority": _priority_label(it.priority, lang),
            "employee": (it.employee_full_name or "").strip() or None,
            "contact_person": (it.contact_person or "").strip() or None,
            "notes": (it.notes or "").strip() or None,
            "form_status": _status_label(getattr(form.status, "value", str(form.status)), lang),
            "hr_decision": _decision_label(form.hr_decision, lang),
            "hr_budget_total": float(form.hr_budget_total) if form.hr_budget_total is not None else None,
            "hr_comment": (form.hr_comment or "").strip() or None,
            "created_at": form.created_at.strftime("%Y-%m-%d %H:%M") if getattr(form, "created_at", None) else "—",
        }

        # build dict keyed by translated headers
        row_out: Dict[str, Any] = {}
        for key in columns_keys:
            row_out[header_of(key, lang)] = values_by_key.get(key)
        rows.append(row_out)

    # build per-form summary sheet
    per_form_summary: List[Dict[str, Any]] = []
    for form_id, meta in sorted(form_meta.items(), key=lambda x: x[0], reverse=True):
        mpk_list = sorted([c for c in form_cost_centers.get(form_id, set()) if c])
        per_form_summary.append(
            summary_row(
                lang,
                form_id=meta["form_id"],
                area=meta["area"],
                mpk_in_items=", ".join(mpk_list) if mpk_list else "—",
                status=meta["status"],
                created_at=meta["created_at"],
                sum_cost=float(form_totals_cost.get(form_id, 0.0)),
                sum_hours=float(form_totals_hours.get(form_id, 0.0)),
                hr_decision=meta["hr_decision"],
                hr_budget_total=meta["hr_budget_total"],
                hr_comment=meta["hr_comment"],
            )
        )

    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")

    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=headers)
        writer.writeheader()
        for r in rows:
            writer.writerow({h: r.get(h) for h in headers})

        writer.writerow({})
        totals_row = {h: "" for h in headers}
        totals_row[headers[0]] = "SUMA" if lang == "pl" else "TOTAL"
        total_cost_h = header_of("total_cost", lang)
        total_hours_h = header_of("total_hours", lang)
        if total_cost_h in headers:
            totals_row[total_cost_h] = f"{totals_cost:.2f}"
        if total_hours_h in headers:
            totals_row[total_hours_h] = f"{totals_hours:.1f}"
        writer.writerow(totals_row)

        data = buf.getvalue().encode("utf-8")
        filename = f"poap_raport_{ts}.csv"
        return StreamingResponse(
            io.BytesIO(data),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    xlsx_bytes = _write_xlsx(lang, rows, headers, totals_cost, totals_hours, per_form_summary)
    filename = f"poap_raport_{ts}.xlsx"
    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
